from services.basic_results import category_area
from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


AREAS = ["Científica", "Humanista", "Artística", "Técnica", "Social"]


def _normalizar_area(area):
    equivalencias = {
        "Técnico-Práctica": "Técnica",
        "Técnico-Práctico": "Técnica",
        "Técnica": "Técnica",
    }
    return equivalencias.get(area, area or "Sin área")


def _ajustar_columnas(hoja):
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[letra].width = min(ancho + 2, 40)


def _estilizar_encabezado(hoja, fila=1):
    relleno = PatternFill("solid", fgColor="0D6EFD")
    for celda in hoja[fila]:
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")


def crear_excel_institucional(results, school_name, periodo_texto):
    """Genera un Excel institucional sin correos ni datos del apoderado."""

    libro = Workbook()
    detalle = libro.active
    detalle.title = "Evaluaciones"

    encabezados = [
        "Año",
        "Fecha",
        "RUT del estudiante",
        "Nombre del estudiante",
        "Curso",
        "Paralelo",
        "Tipo de enseñanza",
        "Área TP",
        "Tipo de test",
        "Área principal",
        "Área secundaria",
        "Puntaje Científica",
        "Puntaje Humanista",
        "Puntaje Artística",
        "Puntaje Técnico-Práctica",
        "Puntaje Social",
        "Ruta sugerida",
        "Estado del registro",
    ]
    detalle.append(encabezados)
    _estilizar_encabezado(detalle)
    detalle.freeze_panes = "A2"
    detalle.auto_filter.ref = f"A1:R{len(results) + 1}"

    for resultado in results:
        fecha = resultado.created_at
        detalle.append([
            fecha.year if fecha else "",
            fecha.strftime("%d-%m-%Y") if fecha else "",
            resultado.student_rut or "",
            resultado.student_name or "",
            resultado.course or "",
            resultado.section or "SIN PARALELO",
            resultado.education_type or "",
            resultado.tp_area or "",
            resultado.test_type or "",
            resultado.observed_main_area or "",
            resultado.observed_secondary_area or "",
            resultado.score_cientifica or 0,
            resultado.score_humanista or 0,
            resultado.score_artistica or 0,
            resultado.score_tecnica or 0,
            resultado.score_social or 0,
            resultado.observed_path or "",
            resultado.status or "activo",
        ])

    _ajustar_columnas(detalle)

    resumen = libro.create_sheet("Resumen")
    resumen.append(["Resumen institucional", school_name])
    resumen.append(["Periodo", periodo_texto])
    resumen.append(["Evaluaciones incluidas", len(results)])
    resumen.append([
        "Estudiantes únicos",
        len({r.student_rut for r in results if r.student_rut}),
    ])
    resumen.append([])

    areas = AREAS + sorted({_normalizar_area(category_area(result, attr))
                            for result in results for attr in ('main_area', 'secondary_area')} - set(AREAS))
    encabezado_resumen = ["Curso", "Evaluaciones", "Estudiantes únicos"]
    encabezado_resumen += [f"Principal: {area}" for area in areas]
    encabezado_resumen += [f"Secundaria: {area}" for area in areas]
    resumen.append(encabezado_resumen)
    _estilizar_encabezado(resumen, 6)

    datos_curso = defaultdict(lambda: {
        "evaluaciones": 0,
        "estudiantes": set(),
        "principales": defaultdict(int),
        "secundarias": defaultdict(int),
    })

    for resultado in results:
        curso = resultado.course or "Sin curso"
        datos = datos_curso[curso]
        datos["evaluaciones"] += 1
        if resultado.student_rut:
            datos["estudiantes"].add(resultado.student_rut)
        datos["principales"][_normalizar_area(category_area(resultado, "main_area"))] += 1
        datos["secundarias"][_normalizar_area(category_area(resultado, "secondary_area"))] += 1

    orden_cursos = {
        "6° Básico": 1,
        "7° Básico": 2,
        "8° Básico": 3,
        "1° Medio": 4,
        "2° Medio": 5,
        "3° Medio": 6,
        "4° Medio": 7,
    }

    for curso in sorted(datos_curso, key=lambda c: (orden_cursos.get(c, 99), c)):
        datos = datos_curso[curso]
        fila = [curso, datos["evaluaciones"], len(datos["estudiantes"])]
        fila += [datos["principales"].get(area, 0) for area in areas]
        fila += [datos["secundarias"].get(area, 0) for area in areas]
        resumen.append(fila)

    resumen.freeze_panes = "A7"
    _ajustar_columnas(resumen)

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida
