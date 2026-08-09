from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _ajustar_columnas(hoja):
    for indice, columna in enumerate(hoja.columns, start=1):
        letra = get_column_letter(indice)
        ancho = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[letra].width = min(ancho + 2, 42)


def _encabezado_excel(hoja, fila):
    relleno = PatternFill("solid", fgColor="123C75")
    for celda in hoja[fila]:
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")


def crear_excel_curso(resultados, datos, school_name, curso_completo, periodo_texto):
    """Genera el Excel del curso respetando los resultados ya filtrados por la ruta."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen del curso"

    resumen.append(["Ficha estadística del curso", ""])
    resumen.merge_cells("A1:E1")
    resumen["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    resumen["A1"].fill = PatternFill("solid", fgColor="123C75")
    resumen["A1"].alignment = Alignment(horizontal="center")
    resumen.append(["Establecimiento", school_name])
    resumen.append(["Curso", curso_completo])
    resumen.append(["Periodo", periodo_texto])
    resumen.append(["Estudiantes únicos", datos["total_estudiantes"]])
    resumen.append(["Evaluaciones incluidas", datos["total_evaluaciones"]])
    resumen.append(["Área principal predominante", datos["principal_predominante"]])
    resumen.append(["Área secundaria predominante", datos["secundaria_predominante"]])
    resumen.append([])
    resumen.append(["Área", "Principal", "% principal", "Secundaria", "% secundaria"])
    _encabezado_excel(resumen, 10)

    for principal, secundaria in zip(datos["principales"], datos["secundarias"]):
        resumen.append([
            principal["area"],
            principal["cantidad"],
            principal["porcentaje"] / 100,
            secundaria["cantidad"],
            secundaria["porcentaje"] / 100,
        ])
    for fila in range(11, 16):
        resumen.cell(fila, 3).number_format = "0.0%"
        resumen.cell(fila, 5).number_format = "0.0%"

    resumen.append([])
    resumen.append(["Interpretación institucional"])
    resumen["A17"].font = Font(bold=True, color="123C75")
    resumen.append([datos["interpretacion"]])
    resumen.merge_cells("A18:E18")
    resumen["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    resumen.row_dimensions[18].height = 62
    resumen.append([])
    resumen.append(["Acciones y talleres sugeridos"])
    resumen["A20"].font = Font(bold=True, color="123C75")
    for accion in datos["acciones"]:
        resumen.append([f"• {accion}"])
        resumen.merge_cells(start_row=resumen.max_row, start_column=1, end_row=resumen.max_row, end_column=5)
        resumen.cell(resumen.max_row, 1).alignment = Alignment(wrap_text=True)

    resumen.freeze_panes = "A10"
    _ajustar_columnas(resumen)

    detalle = libro.create_sheet("Evaluaciones del curso")
    detalle.append([
        "Año", "Fecha", "RUT del estudiante", "Nombre del estudiante",
        "Curso", "Paralelo", "Tipo de enseñanza", "Área TP", "Tipo de test",
        "Área principal", "Área secundaria", "Puntaje Científica",
        "Puntaje Humanista", "Puntaje Artística", "Puntaje Técnico-Práctica",
        "Puntaje Social", "Ruta sugerida", "Estado del registro",
    ])
    _encabezado_excel(detalle, 1)
    detalle.freeze_panes = "A2"

    for resultado in resultados:
        fecha = resultado.created_at
        detalle.append([
            fecha.year if fecha else "",
            fecha.strftime("%d-%m-%Y") if fecha else "",
            resultado.student_rut or "",
            resultado.student_name or "",
            resultado.course or "",
            resultado.section or "SIN PARALELO",
            resultado.education_type or "",
            resultado.tp_area or "",
            resultado.test_type or "",
            resultado.main_area or "",
            resultado.secondary_area or "",
            resultado.score_cientifica or 0,
            resultado.score_humanista or 0,
            resultado.score_artistica or 0,
            resultado.score_tecnica or 0,
            resultado.score_social or 0,
            resultado.suggested_path or "",
            resultado.status or "activo",
        ])

    detalle.auto_filter.ref = f"A1:R{len(resultados) + 1}"
    _ajustar_columnas(detalle)

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida


def crear_pdf_curso(datos, school_name, curso_completo, periodo_texto):
    """Genera una ficha PDF agrupada; no incluye nombres ni RUT."""
    salida = BytesIO()
    documento = SimpleDocTemplate(
        salida,
        pagesize=A4,
        rightMargin=1.7 * cm,
        leftMargin=1.7 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title=f"Ficha estadística {curso_completo}",
    )
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "TituloInstitucional",
        parent=estilos["Title"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        textColor=colors.HexColor("#123C75"),
        alignment=TA_CENTER,
        spaceAfter=12,
    )
    subtitulo = ParagraphStyle(
        "SubtituloInstitucional",
        parent=estilos["Heading2"],
        textColor=colors.HexColor("#123C75"),
        fontSize=12,
        leading=15,
        spaceBefore=10,
        spaceAfter=7,
    )
    normal = ParagraphStyle(
        "NormalInstitucional",
        parent=estilos["BodyText"],
        fontSize=9.5,
        leading=13,
    )
    pequeno = ParagraphStyle(
        "PequenoInstitucional",
        parent=normal,
        fontSize=8,
        textColor=colors.HexColor("#555555"),
    )

    elementos = [
        Paragraph("Sistema Integral de Orientación Educacional", titulo),
        Paragraph("Ficha estadística agregada del curso", subtitulo),
    ]
    identificacion = Table([
        ["Establecimiento", school_name],
        ["Curso", curso_completo],
        ["Periodo", periodo_texto],
    ], colWidths=[4.2 * cm, 12.2 * cm])
    identificacion.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF1FA")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#123C75")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C6D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elementos.extend([identificacion, Spacer(1, 10)])

    indicadores = Table([
        ["Estudiantes únicos", "Evaluaciones", "Principal predominante", "Secundaria predominante"],
        [str(datos["total_estudiantes"]), str(datos["total_evaluaciones"]),
         datos["principal_predominante"], datos["secundaria_predominante"]],
    ], colWidths=[3.6 * cm, 3.4 * cm, 4.8 * cm, 4.8 * cm])
    indicadores.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123C75")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.3),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C6D9")),
        ("PADDING", (0, 0), (-1, -1), 7),
    ]))
    elementos.extend([indicadores, Paragraph("Distribución de áreas", subtitulo)])

    filas = [["Área", "Principal", "%", "Secundaria", "%"]]
    for principal, secundaria in zip(datos["principales"], datos["secundarias"]):
        filas.append([
            principal["area"], str(principal["cantidad"]), f'{principal["porcentaje"]}%',
            str(secundaria["cantidad"]), f'{secundaria["porcentaje"]}%',
        ])
    tabla_areas = Table(filas, colWidths=[5.2 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm, 2.8 * cm])
    tabla_areas.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#123C75")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.7),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C6D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elementos.extend([tabla_areas, Paragraph("Interpretación institucional", subtitulo),
                       Paragraph(datos["interpretacion"], normal)])

    if datos["acciones"]:
        acciones = [Paragraph(f"• {accion}", normal) for accion in datos["acciones"]]
    else:
        acciones = [Paragraph(
            "Se recomienda ampliar las oportunidades de exploración en las distintas áreas.",
            normal,
        )]
    elementos.append(KeepTogether([Paragraph("Acciones y talleres sugeridos", subtitulo)] + acciones))
    elementos.extend([
        Spacer(1, 10),
        Paragraph(
            "La información presenta tendencias grupales. Apoya la planificación del "
            "establecimiento y no reemplaza el análisis profesional del orientador.",
            pequeno,
        ),
    ])

    def pie_pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#667085"))
        canvas.drawString(1.7 * cm, 0.8 * cm, "Sistema Integral de Orientación Educacional")
        canvas.drawRightString(A4[0] - 1.7 * cm, 0.8 * cm, f"Página {doc.page}")
        canvas.restoreState()

    documento.build(elementos, onFirstPage=pie_pagina, onLaterPages=pie_pagina)
    salida.seek(0)
    return salida
